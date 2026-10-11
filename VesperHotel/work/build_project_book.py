from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\User\Documents\Codex\2026-10-08\new-chat\outputs\VesperHotel")
DOCS = ROOT / "docs"
DOCS.mkdir(exist_ok=True)
DOCX = DOCS / "VESPER-project-book-he.docx"
MD = DOCS / "VESPER-project-book-he.md"


def set_rtl(paragraph, rtl=True):
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if rtl and bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


class Book:
    def __init__(self):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11)
        sec.top_margin = Inches(0.72)
        sec.bottom_margin = Inches(0.68)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)
        self.md = []
        self._styles()
        self._header_footer()

    def _styles(self):
        styles = self.doc.styles
        for name, size, bold in [("Normal", 10.5, False), ("Title", 30, True), ("Subtitle", 13, False),
                                 ("Heading 1", 20, True), ("Heading 2", 15, True), ("Heading 3", 12, True)]:
            style = styles[name]
            style.font.name = "Arial"
            style.font.size = Pt(size)
            style.font.bold = bold
            style.font.color.rgb = RGBColor(0, 0, 0)
            style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
            style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
            style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        styles["Normal"].paragraph_format.space_after = Pt(6)
        styles["Normal"].paragraph_format.line_spacing = 1.15
        styles["Heading 1"].paragraph_format.space_before = Pt(16)
        styles["Heading 1"].paragraph_format.space_after = Pt(8)
        styles["Heading 1"].paragraph_format.keep_with_next = True
        styles["Heading 2"].paragraph_format.space_before = Pt(12)
        styles["Heading 2"].paragraph_format.space_after = Pt(5)
        styles["Heading 2"].paragraph_format.keep_with_next = True
        styles["Heading 3"].paragraph_format.space_before = Pt(9)
        styles["Heading 3"].paragraph_format.space_after = Pt(4)
        styles["Heading 3"].paragraph_format.keep_with_next = True

    def _header_footer(self):
        sec = self.doc.sections[0]
        hp = sec.header.paragraphs[0]
        hp.text = "VESPER  ספר הפרויקט"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hp.runs[0].font.name = "Arial"
        hp.runs[0].font.size = Pt(8)
        hp.runs[0].font.color.rgb = RGBColor(80, 91, 82)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fp.add_run("עמוד ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        run._r.addnext(fld)
        fp.add_run("  |  גרסה 1.0  |  אוקטובר 2026")
        for r in fp.runs:
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(100, 100, 100)

    def title(self, title, subtitle):
        p = self.doc.add_paragraph(style="Title")
        set_rtl(p)
        p.paragraph_format.space_before = Pt(90)
        p.add_run(title)
        p2 = self.doc.add_paragraph(style="Subtitle")
        set_rtl(p2)
        p2.add_run(subtitle)
        p3 = self.doc.add_paragraph()
        set_rtl(p3)
        p3.paragraph_format.space_before = Pt(28)
        r = p3.add_run("מדריך לימודי מלא למערכת המלון  Backend  Frontend  Database  Messaging  Infrastructure")
        r.bold = True
        r.font.size = Pt(12)
        p4 = self.doc.add_paragraph()
        set_rtl(p4)
        p4.paragraph_format.space_before = Pt(180)
        p4.add_run("מבוסס על הקוד בפועל בתיקיית VesperHotel\nגרסה 1.0  11 באוקטובר 2026")
        self.doc.add_page_break()
        self.md += [f"# {title}", "", subtitle, "", "גרסה 1.0 — 11 באוקטובר 2026", ""]

    def h1(self, text, page=True):
        if page and len(self.doc.paragraphs) > 6:
            self.doc.add_page_break()
        p = self.doc.add_paragraph(style="Heading 1")
        set_rtl(p)
        p.add_run(text)
        self.md += [f"# {text}", ""]

    def h2(self, text):
        p = self.doc.add_paragraph(style="Heading 2")
        set_rtl(p)
        p.add_run(text)
        self.md += [f"## {text}", ""]

    def h3(self, text):
        p = self.doc.add_paragraph(style="Heading 3")
        set_rtl(p)
        p.add_run(text)
        self.md += [f"### {text}", ""]

    def p(self, text, bold_lead=None):
        p = self.doc.add_paragraph()
        set_rtl(p)
        if bold_lead and text.startswith(bold_lead):
            p.add_run(bold_lead).bold = True
            p.add_run(text[len(bold_lead):])
        else:
            p.add_run(text)
        self.md += [text, ""]

    def bullets(self, items, numbered=False):
        for i, item in enumerate(items, 1):
            p = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            set_rtl(p)
            p.add_run(item)
            self.md.append((f"{i}. " if numbered else "- ") + item)
        self.md.append("")

    def code(self, code):
        for line in code.strip("\n").splitlines():
            p = self.doc.add_paragraph()
            set_rtl(p, False)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(35, 45, 39)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        self.md += ["```", code.strip("\n"), "```", ""]

    def table(self, headers, rows, widths=None, font_size=8.5):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.autofit = True
        hdr = table.rows[0]
        set_repeat_table_header(hdr)
        for i, h in enumerate(headers):
            cell = hdr.cells[i]
            set_cell_shading(cell, "263A31")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            set_rtl(p)
            r = p.add_run(str(h))
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.name = "Arial"
            r.font.size = Pt(font_size)
            if widths:
                cell.width = Inches(widths[i])
        for ri, row in enumerate(rows):
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cell = cells[i]
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ri % 2:
                    set_cell_shading(cell, "F2F5F1")
                p = cell.paragraphs[0]
                set_rtl(p)
                r = p.add_run(str(value))
                r.font.name = "Arial"
                r.font.size = Pt(font_size)
                if widths:
                    cell.width = Inches(widths[i])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        self.md.append("| " + " | ".join(headers) + " |")
        self.md.append("| " + " | ".join(["---"] * len(headers)) + " |")
        for row in rows:
            self.md.append("| " + " | ".join(str(v).replace("\n", "<br>").replace("|", "\\|") for v in row) + " |")
        self.md.append("")

    def save(self):
        self.doc.core_properties.title = "ספר הפרויקט VESPER"
        self.doc.core_properties.subject = "מדריך מלא לארכיטקטורה ולקוד של מערכת המלון"
        self.doc.core_properties.author = "OpenAI Codex"
        self.doc.save(DOCX)
        MD.write_text("\n".join(self.md), encoding="utf-8")


b = Book()
b.title("ספר הפרויקט VESPER", "מדריך מלא להבנת המערכת מן הדפדפן ועד מסד הנתונים")

b.h1("איך להשתמש בספר", page=False)
b.p("הספר נכתב עבורך כמפתח C# שמכיר מחלקות Interfaces LINQ ו async await אך עדיין לא למד Frontend. המטרה היא לא רק להגיד איפה כל קובץ נמצא אלא להסביר למה הוא קיים ואיך המידע עובר בין החלקים. מומלץ לקרוא תחילה את פרקי הארכיטקטורה והזרימות ורק אחר כך להשתמש בנספחי הפונקציות והאובייקטים כחומר עזר בזמן קריאת הקוד.")
b.p("המסקנה המרכזית היא שהמערכת בנויה סביב מסד הנתונים כמקור האמת. האתר מציג ושולח בקשות. ה API מאמת אותן ומפעיל שירותי C#. השירותים משנים את MySQL או SQLite בתוך טרנזקציה. אירועים נרשמים באותה טרנזקציה לטבלת Outbox. Worker נפרד מפרסם אותם ל Kafka מעבד אותם פעם אחת בעזרת Inbox יוצר התראות ושולח נתוני תפעול ל Elasticsearch. Kibana קוראת מ Elasticsearch ומאפשרת לצפות באירועים.")
b.h2("מפת הקריאה")
b.bullets([
    "פרקים 1 עד 4 מסבירים מה המערכת עושה ואיך היא מחולקת.",
    "פרקים 5 עד 8 מלמדים את ה Frontend מהיסוד ואת החיבור ל API.",
    "פרקים 9 עד 14 מסבירים Backend SQL Authentication Validation Transactions Outbox Kafka ו Elastic.",
    "פרקים 15 עד 18 מציגים זרימות מקצה לקצה migrations Docker ובדיקות.",
    "הנספחים הם קטלוג מלא של הפונקציות האובייקטים והנתיבים במערכת."
])
b.h2("תוכן עניינים")
b.bullets([
    "1 מטרת המערכת והגבולות העסקיים", "2 הארכיטקטורה הכללית", "3 מבנה התיקיות", "4 מילון מושגים",
    "5 יסודות Frontend", "6 React בפרויקט", "7 שכבת API בצד הדפדפן", "8 מסכי האתר והאדמין",
    "9 Backend ב ASP.NET Core", "10 Authentication והרשאות", "11 Validation ושגיאות", "12 מסד הנתונים והיחסים",
    "13 טרנזקציות מקביליות ו Idempotency", "14 Outbox Inbox Kafka Elasticsearch ו Kibana",
    "15 API מזג האוויר", "16 Migrations ו Seed", "17 הזרימות המלאות", "18 Docker בדיקות ותחזוקה",
    "נספח א כל פונקציות C#", "נספח ב כל פונקציות Frontend", "נספח ג כל האובייקטים", "נספח ד כל נתיבי ה API"
])

b.h1("1 מטרת המערכת והגבולות העסקיים")
b.p("VESPER היא מערכת הזמנות למלון בדיוני בלאס וגאס. אותו קטלוג מכיל שלושה סוגי משאבים: חדרים פעילויות ותחבורה. כל סוג נראה דומה למשתמש אך משתמש במודל מלאי שונה. חדר או רכב מקבלים יחידה פיזית. פעילות מקבלת Session עם שעה קיבולת ומחיר. ההבחנה הזאת היא בסיס התכנון של כל תהליך ההזמנה.")
b.table(["תחום", "החלטה עסקית", "השפעה בקוד"], [
    ["חדר", "הזמנה של 1 עד 30 לילות כניסה ב 15:00 יציאה ב 11:00", "BookingRules.Period מחשבת UTC ו BookingService מקצה Unit"],
    ["פעילות", "הלקוח בוחר מועד מוגדר ומספר מושבים", "ActivitySession מחזיק Capacity והקצאה סופרת Guests"],
    ["תחבורה", "מחיר לפי שעות ומרווח תפעולי של 30 דקות", "InventoryEndUtc מאוחר מ EndUtc בשלושים דקות"],
    ["Hold", "מלאי נשמר לעשר דקות", "Booking במצב Pending עם ExpiresAtUtc"],
    ["מחיר", "המחיר ננעל בעת יצירת ה Hold", "TotalCents ו PriceLinesJson נשמרים בתוך Booking"],
    ["תשלום", "אין סליקה באתר והתשלום במלון", "אין אובייקט Payment ואין שמירת כרטיס"],
    ["ביטול", "48 שעות לחדר ו 24 שעות לפעילות או תחבורה", "CancellationDeadlineUtc נבדק ב TransitionAsync"],
    ["מנהל", "יכול לבטל באיחור רק עם סיבה", "Authorize Role Admin ו AuditEntry"],
])
b.p("המערכת מונעת Overbooking גם כאשר בקשות מגיעות במקביל. היא גם שומרת את ההיסטוריה: משאב מושבת במקום להימחק והזמנה שומרת צילום של שם המשאב התמונה המחיר היחידה ותנאי הביטול כפי שהיו בזמן ההזמנה.")

b.h1("2 הארכיטקטורה הכללית")
b.p("הארכיטקטורה היא אפליקציית Web מודולרית עם תהליך רקע. React רץ בדפדפן. ASP.NET Core מספק API. Entity Framework Core מחבר את השירותים למסד הנתונים. במצב Docker מסד הנתונים הוא MySQL והעיבוד האסינכרוני משתמש ב Kafka. Elasticsearch מחזיק אינדקס אירועים ו Kibana מציגה אותו. במצב פיתוח קל SQLite מחליף את MySQL ו LocalHotelWorker מעבד את ה Outbox בלי Kafka.")
b.code("""
Browser React
    | HTTPS JSON
    v
ASP.NET Core API  ->  Open Meteo Forecast API
    | Entity Framework Core
    v
MySQL or SQLite  <- source of truth
    | Outbox rows
    v
Hotel Worker -> Kafka -> Inbox -> Notifications
    |                         \\-> local delivery files
    \\-> Elasticsearch -> Kibana
""")
b.h2("חלוקת האחריות")
b.table(["שכבה", "אחריות", "מה היא אינה עושה"], [
    ["Frontend", "תצוגה ניווט טפסים מצב זמני ושליחת בקשות", "אינו קובע מחיר או זמינות סופית"],
    ["Controllers", "תרגום HTTP לקריאה לשירות והחזרת status code", "אינם מחזיקים כללים עסקיים גדולים"],
    ["Services", "Authentication הזמנות קטלוג ניהול ומזג אוויר", "אינם מציירים UI"],
    ["Domain", "אובייקטי ליבה וחוקי זמן מחיר וחפיפה", "אינו תלוי ב React או Kafka"],
    ["Infrastructure", "EF Core טרנזקציות Identity Workers ושירותים חיצוניים", "אינו קובע את מבנה המסכים"],
    ["Database", "מקור אמת מתמשך נעילות אינדקסים ויחסים", "אינו תחליף ל Validation עסקי"],
    ["Worker", "פקיעת הזמנות פרסום אירועים התראות ואינדוקס", "אינו עונה לבקשות דפדפן"],
])
b.h2("למה אין שכבות רבות יותר")
b.p("הקוד נשמר קרוב לרמה של פרויקט C# לימודי: מחלקות שירות ברורות DbContext ו LINQ בלי מסגרת CQRS גדולה בלי MediatR ובלי עשרות Interfaces שאין להם כרגע יותר ממימוש אחד. ההפרדה עדיין חזקה במקום שבו היא חשובה: Domain אינו מכיר HTTP Controllers דקים ושינויי מלאי נעשים דרך שירותים וטרנזקציות.")
b.h2("Interfaces בפרויקט")
b.p("ב Backend אין Interfaces מותאמים אישית כגון IBookingService משום שלכל שירות יש כרגע מימוש אחד בלבד והחלפה בין מימושים אינה דרישה עסקית. Dependency Injection עדיין מזריקה את המחלקות עצמן. Interfaces כן מופיעים דרך הספריות כגון IServiceCollection IConfiguration IMemoryCache IProducer ו IConsumer. ב Frontend המילה interface משמשת בעיקר לתיאור צורת JSON בזמן build כגון CatalogItem Booking ו User. אלה חוזים של TypeScript ואינם אובייקטים שקיימים בזמן ריצה.")

b.h1("3 מבנה התיקיות")
b.table(["תיקייה", "תוכן", "נקודת התחלה מומלצת"], [
    ["src Hotel.Domain", "ישויות הליבה Enums וכללי הזמנה טהורים", "Models.cs ואז BookingRules.cs"],
    ["src Hotel.Infrastructure", "DbContext DTOs Services Workers ו Migrations של SQLite", "HotelDbContext.cs ואז BookingService.cs"],
    ["src Hotel.Api", "Program middleware ו Controllers", "Program.cs ואז Controllers"],
    ["src Hotel.Worker", "תהליך הרקע העצמאי", "Program.cs ו HotelWorker.cs"],
    ["src Hotel.Migrations.MySql", "Migrations מותאמות ל MySQL", "קובצי migration לפי תאריך"],
    ["frontend hotel-web src", "React TypeScript routes hooks pages styles", "main.tsx App.tsx api.ts"],
    ["frontend hotel-web public", "תמונות פונטים וקבצים סטטיים", "images rooms"],
    ["infrastructure", "Docker Compose Dockerfiles ו nginx", "docker-compose.yml"],
    ["scripts", "אתחול סודות הפעלה עצירה וסטטוס", "start-docker.ps1 ו common.ps1"],
    ["tests", "בדיקת API מקצה לקצה", "api-smoke.mjs"],
    ["docs", "מדריכי לימוד אימות וספר הפרויקט", "הקובץ הנוכחי"],
])
b.h2("התלויות בין פרויקטי C#")
b.bullets([
    "Hotel.Domain עומד בבסיס ואינו מפנה לפרויקטים האחרים.",
    "Hotel.Infrastructure מפנה ל Domain ומממש גישה לנתונים ושירותים.",
    "Hotel.Api מפנה ל Infrastructure ול Domain ומארח את HTTP API.",
    "Hotel.Worker מפנה ל Infrastructure ומריץ את HotelWorker.",
    "Hotel.Migrations.MySql מפנה ל Infrastructure ומכיל snapshots ומיגרציות עבור MySQL."
])

b.h1("4 מילון מושגים")
b.table(["מושג", "הסבר פשוט"], [
    ["API", "חוזה של כתובות HTTP שה Frontend יכול לקרוא או לשנות דרכן מידע"],
    ["Endpoint", "כתובת ופעולה מסוימת כגון GET api catalog או POST api bookings"],
    ["DTO", "אובייקט להעברת נתונים דרך ה API שאינו בהכרח ישות מסד נתונים"],
    ["Entity", "אובייקט שממופה בדרך כלל לשורה בטבלת SQL"],
    ["DbContext", "המחלקה שדרכה EF Core בונה שאילתות ושומר שינויים"],
    ["Migration", "תיאור מדורג של שינוי מבנה מסד הנתונים"],
    ["JWT", "Access token חתום שמוכיח מי המשתמש ומה ה Role שלו"],
    ["Refresh token", "סוד ארוך יותר שמאפשר לקבל Access token חדש"],
    ["Claim", "עובדה חתומה בתוך JWT כגון UserId Email או Role"],
    ["Transaction", "קבוצת שינויי SQL שנשמרים כולם או לא נשמרים כלל"],
    ["Idempotency", "אותה בקשה חוזרת אינה יוצרת פעולה כפולה"],
    ["Optimistic concurrency", "Version מזהה שמישהו ערך את הרשומה מאז שקראנו אותה"],
    ["Outbox", "טבלת אירועים שנכתבת באותה טרנזקציה עם השינוי העסקי"],
    ["Inbox", "טבלה שמוכיחה שאירוע כבר עובד ומונעת עיבוד כפול"],
    ["Kafka", "Broker שמעביר אירועים בין Producer ל Consumer"],
    ["Elasticsearch", "מסד חיפוש למסמכים ואירועים שאינו מקור האמת העסקי"],
    ["Kibana", "ממשק חיפוש ו Dashboard מעל Elasticsearch"],
    ["Hook", "פונקציית React כגון useState או useEffect שמוסיפה התנהגות ל Component"],
    ["Props", "קלט שמועבר מ Component הורה ל Component ילד"],
    ["State", "מידע זמני של המסך שכאשר הוא משתנה React מצייר מחדש"],
])

b.h1("5 יסודות Frontend למפתח C#")
b.h2("מה קורה בדפדפן")
b.p("הדפדפן מקבל index.html קטן. הקובץ main.tsx מפעיל את React ומחבר את App לאלמנט root. React יוצר עץ Components. כל Component הוא פונקציה שמחזירה JSX. JSX נראה כמו HTML אך הוא תחביר TypeScript שמומר לקריאות JavaScript בזמן build.")
b.code("""
function Greeting({ name }: { name: string }) {
  return <h1>Hello {name}</h1>;
}
""")
b.p("בהשוואה ל C# אפשר לחשוב על Props כעל פרמטרים לפונקציה ועל JSX כתוצאה המתארת מה צריך להופיע במסך. React מפעיל שוב את הפונקציה כאשר State או Props משתנים ומשווה את התוצאה לעץ הקודם.")
b.h2("JavaScript TypeScript ו TSX")
b.bullets([
    "JavaScript היא השפה שהדפדפן מפעיל.",
    "TypeScript מוסיפה טיפוסים בזמן הפיתוח. הטיפוסים נעלמים לאחר build ולכן גם השרת חייב לבצע Validation.",
    "קובץ ts מכיל TypeScript רגיל. קובץ tsx מאפשר גם JSX.",
    "interface ב TypeScript מתאר צורה של אובייקט JSON. הוא דומה למחלקת DTO ב C# אך אין לו מימוש בזמן ריצה.",
    "Vite בונה את הקוד לקובצי JavaScript ו CSS שהדפדפן יכול לטעון."
])
b.h2("HTML ו CSS בפרויקט")
b.p("HTML מתאר משמעות ומבנה: header nav main section form input button table ו dialog. CSS קובע צבעים מרווחים גרידים התאמה למובייל ואנימציות. className ב JSX הופך ל class רגיל ב HTML. הקובץ styles.css מרכז את העיצוב. Media queries בסוף הקובץ משנות את הפריסה במסכים קטנים.")
b.table(["מושג Frontend", "מקבילה שימושית ב C#", "דוגמה בפרויקט"], [
    ["Component", "מתודה שמקבלת DTO ומחזירה View", "ResourceCard"],
    ["Props", "פרמטרים", "item ו index"],
    ["State", "שדות זמניים של ViewModel", "dates guests busy error"],
    ["useEffect", "קוד שרץ לאחר שינוי תלות", "טעינת API או הוספת keyboard listener"],
    ["Context", "שירות משותף בטווח עץ UI", "AuthContext"],
    ["Hook מותאם", "מתודת עזר עם State", "useApi"],
    ["Route", "מיפוי URL למסך", "rooms אל CatalogPage"],
])

b.h1("6 React בפרויקט")
b.h2("נקודת הכניסה וה Routing")
b.p("main.tsx יוצר BrowserRouter עוטף את האתר ב AuthProvider ומרנדר את App. BrowserRouter קורא את כתובת הדפדפן ללא טעינת עמוד מלאה. App.tsx מגדיר Routes. PublicLayout משותף למסכים הציבוריים. Guard עוטף מסכים שמחייבים משתמש או Admin.")
b.p("כאשר כתובת משתנה React Router בוחר Component חדש. Outlet הוא המקום שבו הילד של Route מצויר. Navigate מבצע הפניה. Link משנה כתובת ללא refresh מלא ולכן ה State הגלובלי נשמר.")
b.h2("State ו Hooks")
b.p("useState מחזיר ערך ופונקציית עדכון. לדוגמה const busy setBusy שומר אם בקשה פעילה. קריאה ל setBusy true גורמת לציור מחדש שבו הכפתור מושבת. useEffect מתאים לפעולות צד כגון fetch או listener. בכל Effect שמתחיל פעולה מתמשכת יש cleanup כדי למנוע עדכון לאחר שה Component ירד מהמסך.")
b.h2("ה Hook useApi")
b.p("useApi מקבל path. הוא שומר data error loading ו revision. Effect קורא ל api כאשר path או revision משתנים. refresh מגדיל revision וכך מפעיל את הקריאה מחדש. המשתנה active מונע מתגובה מאוחרת לעדכן Component שכבר אינו מוצג.")
b.h2("טפסים")
b.p("רוב השדות הם Controlled Inputs: value מגיע מ State ו onChange מעדכן State. onSubmit מבטל את שליחת הטופס הרגילה באמצעות preventDefault ואז קורא ל send. Validation בסיסי כמו required min max ו maxLength עוזר למשתמש אך אינו גבול אבטחה. ה Backend מאמת שוב את הכול.")
b.h2("נגישות")
b.p("האתר משתמש ב label לכל input ב aria-label לכפתורי אייקון ב role dialog ל Modal וב role status או alert להודעות. מקש Escape סוגר Modal וגלריה. כותרות ו landmarks כמו header nav main ו footer מאפשרים לקורא מסך להבין את העמוד.")

b.h1("7 שכבת API בצד הדפדפן")
b.p("api.ts הוא השער היחיד כמעט לכל תקשורת HTTP. הוא מוסיף Authorization אם יש Access token מוסיף credentials include כדי שה Cookie של Refresh יישלח ומתרגם שגיאות JSON ל ApiError. אם בקשה מוגנת מחזירה 401 הוא מנסה פעם אחת לבצע refresh ואז משחזר את הבקשה המקורית.")
b.code("""
const response = await fetch(`/api${path}`, {
  ...options,
  headers,
  credentials: 'include'
});
""")
b.h2("למה Access token נשמר בזיכרון")
b.p("ה Access token אינו נשמר ב localStorage. הוא נמצא במשתנה JavaScript ונעלם ברענון. בזמן טעינת האתר AuthProvider שולח את Refresh cookie לשרת ומקבל Access token חדש. ה Refresh token עצמו נמצא ב Cookie מסוג HttpOnly ולכן JavaScript אינו יכול לקרוא אותו. כך מצמצמים את הנזק אם קוד זדוני מצליח לקרוא אחסון דפדפן.")
b.h2("עיצוב מצב טעינה ושגיאה")
b.p("כל מסך שמבקש נתונים מציג אחת מארבע אפשרויות: Loading בזמן המתנה ErrorBox במקרה כשל תוכן במקרה הצלחה ו Empty כאשר התוצאה תקינה אך ריקה. ההפרדה הזאת חשובה כי null אינו אומר תמיד שגיאה.")

b.h1("8 מסכי האתר וממשק הניהול")
b.table(["Component", "כתובת", "אחריות"], [
    ["HomePage", "/", "Hero חיפוש חדרים מזג אוויר וחדרים מובילים"],
    ["CatalogPage", "/rooms /experiences /journeys", "טעינת קטלוג סינון וסידור"],
    ["DetailPage", "/discover/:id", "פרטי משאב גלריית חדר ופאנל הזמנה"],
    ["AuthPage", "/login /register", "כניסה או הרשמה של אורח"],
    ["AdminLoginPage", "/admin/login", "כניסת מנהל ובדיקת Role"],
    ["BookingPage", "/bookings/:id", "בדיקת Hold טיימר אישור וביטול"],
    ["AccountLayout", "/account", "מעטפת לחשבון הזמנות פרופיל והתראות"],
    ["AdminDashboard", "/admin", "מדדים תפוסה ערך הזמנות ופעילות קרובה"],
    ["AdminBookings", "/admin/bookings", "חיפוש סינון צפייה וביטול מנהלי"],
    ["AdminResources", "/admin/resources", "משאבים יחידות Sessions מחיר וגלריות"],
    ["AdminUsers", "/admin/users", "חיפוש משתמשים"],
    ["AdminAudit", "/admin/audit", "הצגת היסטוריית שינויים"],
    ["AdminSystem", "/admin/system", "בריאות Outbox כשלי אירועים והתראות"],
])
b.h2("זרימת BookingPanel")
b.bullets([
    "המשתמש בוחר תאריכים אורחים או Session.",
    "המסך קורא GET availability ומציג מחיר וזמינות מהשרת.",
    "בלחיצה Reserve נוצר Idempotency Key אקראי ונשלחת POST bookings.",
    "אם אין משתמש Guard או טיפול הניווט שולחים אותו למסך כניסה עם returnTo.",
    "לאחר יצירת Hold עוברים לעמוד ההזמנה שבו מוצג HoldTimer.",
    "Confirm שולח POST bookings id confirm. Cancel שולח reason ל cancel."
])
b.h2("גלריית החדר")
b.p("RoomGallery מחברת את imageUrl ואת galleryImages מסירה כפילויות ומציגה שלוש תמונות בפריסה. לחיצה פותחת lightbox. State בשם selected מחזיק את אינדקס התמונה. Effect מוסיף טיפול ב Escape ArrowLeft ו ArrowRight ונועל גלילה של body. כאשר סוגרים את הגלריה ה Effect מסיר את המאזין ומשחזר גלילה.")

b.h1("9 Backend ב ASP.NET Core")
b.h2("Program.cs")
b.p("Program.cs הוא Composition Root. הוא קורא Configuration רושם שירותים ב Dependency Injection מגדיר Controllers JSON CORS Rate Limiting Swagger Authentication ו Authorization בונה את האפליקציה מרכיב Middleware ממפה health ו Controllers מפעיל migrations ו seed ואז מתחיל להאזין.")
b.h2("Dependency Injection")
b.p("AddHotel רושמת HotelDbContext ושירותים כ Scoped. בכל בקשת HTTP מתקבל מופע חדש של השירות וה DbContext. WeatherService נרשם כ Typed HttpClient. Worker משתמש IServiceScopeFactory כדי לפתוח Scope חדש בכל פעולת רקע וכך אינו שומר DbContext לאורך זמן.")
b.h2("Controllers")
b.p("Controller מקבל DTO מ JSON או query string ASP.NET מבצע model binding ו DataAnnotations Validation ואז המתודה קוראת לשירות. ActionResult מאפשר להחזיר גם אובייקט וגם status code כגון 201 או 503. Authorize דורש JWT תקין ו Authorize Roles Admin דורש Claim תפקיד Admin.")
b.h2("Middleware שגיאות")
b.p("ApiErrors עוטף כל בקשה. הוא מוסיף Correlation ID וכותרת nosniff מודד זמן ומתרגם HotelException לשגיאת API צפויה. DbUpdateConcurrencyException הופכת ל 409 version conflict. שגיאה לא צפויה נרשמת ל log ומוחזרת כ 500 בלי פרטים פנימיים.")

b.h1("10 Authentication והרשאות")
b.h2("מהי Authentication")
b.p("Authentication היא תשובה לשאלה מי המשתמש. Authorization היא תשובה לשאלה מה מותר לו לעשות. בפרויקט ASP.NET Identity מנהל משתמשים סיסמאות Roles נעילה וטוקנים. JWT Bearer Authentication קורא את הכותרת Authorization Bearer מאמת חתימה תוקף Issuer ו Audience ויוצר ClaimsPrincipal.")
b.h2("הרשמה וכניסה")
b.bullets([
    "RegisterAsync יוצר HotelUser בעזרת UserManager. Identity מבצע hashing לסיסמה ואינו שומר אותה כטקסט.",
    "המשתמש מקבל Role בשם Guest.",
    "LoginAsync מוצא משתמש ומפעיל CheckPasswordSignInAsync עם lockoutOnFailure.",
    "לאחר הצלחה IssueAsync יוצר JWT ל 15 דקות ו Refresh token אקראי ל 7 ימים.",
    "במסד נשמר רק SHA256 hash של Refresh token. הערך המקורי נשמר ב HttpOnly cookie.",
    "RefreshAsync מבטל את ה Refresh token הישן ומנפיק זוג חדש. זה נקרא token rotation.",
    "LogoutAsync מסמן את ה session כמבוטל ומוחק את ה Cookie."
])
b.h2("Claims ב JWT")
b.table(["Claim", "ערך", "שימוש"], [
    ["NameIdentifier", "User Id", "איתור ההזמנות והפרופיל של המשתמש"],
    ["Email", "כתובת דואר", "Audit ותצוגה"],
    ["Role", "Guest או Admin", "Authorize Roles Admin ו User.IsInRole"],
    ["stamp", "SecurityStamp של Identity", "פסילת Access token לאחר שינוי אבטחה"],
])
b.h2("דוגמת שימוש ב C#")
b.code("""
[Authorize(Roles = "Admin")]
[HttpGet("resources")]
public Task<List<CatalogItem>> Resources()
    => catalog.ListAsync(null, true);

var userId = User.FindFirstValue(ClaimTypes.NameIdentifier);
var isAdmin = User.IsInRole("Admin");
""")
b.p("OnTokenValidated ב ServiceRegistration משווה את Claim stamp ל SecurityStamp העדכני במסד. אם הסיסמה או מצב האבטחה שינו את ה stamp הטוקן הישן נפסל גם אם טרם הגיע זמן התפוגה.")
b.h2("אבטחת סיסמאות ואיפוס")
b.p("מדיניות הסיסמה דורשת לפחות 10 תווים אות גדולה אות קטנה ספרה ותו שאינו אלפאנומרי. לאחר חמש כניסות כושלות המשתמש ננעל ל 15 דקות. ForgotAsync תמיד מחזירה הודעה כללית כדי לא לחשוף אם אימייל קיים. אם הוא קיים נוצר token חתום לשעה וקישור האיפוס נכתב ל Outbox. ResetAsync מבצעת את האיפוס דרך Identity ומבטלת את כל RefreshSessions הפעילים.")

b.h1("11 Validation ושגיאות")
b.p("Validation מתבצע בכמה שכבות כי לכל שכבה תפקיד אחר. HTML עוזר למשתמש. TypeScript מצמצם טעויות מפתח. DataAnnotations בודקות מבנה קלט בצד השרת. שירותי C# בודקים כללים עסקיים. מסד הנתונים אוכף uniqueness foreign keys ו concurrency.")
b.table(["שכבה", "דוגמאות", "מה קורה בכשל"], [
    ["HTML", "required min max maxLength", "הדפדפן מציג הודעה לפני שליחה"],
    ["TypeScript", "ResourceKind union ו interfaces", "שגיאת build בזמן פיתוח"],
    ["DataAnnotations", "Required EmailAddress Range MaxLength", "400 validation_failed עם errors"],
    ["Domain", "תאריך חדר 1 עד 30 לילות מחיר עד שתי ספרות", "HotelException עם code ברור"],
    ["Service", "OfferVersion סוג משאב Route גלריה ו Session", "400 או 409 לפי הבעיה"],
    ["Database", "FK unique indexes concurrency tokens", "409 conflict דרך ApiErrors"],
])
b.h2("HotelException")
b.p("HotelException מכילה Status ו Code בנוסף ל Message. השירות זורק לדוגמה new HotelException 409 unavailable. Middleware תופס אותה ומחזיר JSON אחיד עם status code message ו traceId. כך השירות אינו תלוי ב MVC ועדיין שולט במשמעות השגיאה.")
b.h2("ולידציות מרכזיות")
b.bullets([
    "מחיר בין 0 ל 1,000,000 ועד שתי ספרות אחרי הנקודה.",
    "חדר מקבל תאריכי YYYY-MM-DD בלבד ושהייה של 1 עד 30 לילות.",
    "פעילות או תחבורה מקבלות ISO timestamps וסיום לאחר התחלה בתוך שבעה ימים.",
    "כל הזמנה מתחילה בעתיד ומספר האורחים בין 1 ל MaxGuests.",
    "תחבורה מחייבת PickupLocation ו DropoffLocation.",
    "Hold מחייב Idempotency-Key באורך 1 עד 120.",
    "Image URL חייב להתחיל ב /images/ או להיות HTTPS. גלריה מוגבלת ל 12 תמונות.",
    "Session שמכיל הזמנות אינו ניתן להזזה ואי אפשר להקטין Capacity מתחת למספר השמורים.",
    "Unit עם הזמנה פעילה אינו ניתן להשבתה.",
    "ExpectedVersion ו ExpectedOfferVersion מונעים שמירה על נתון מיושן."
])

b.h1("12 מסד הנתונים והיחסים")
b.p("Entity Framework Core ממפה מחלקות ל SQL. DbSet הוא נקודת הכניסה לטבלה. LINQ מתורגם לשאילתת SQL. SaveChangesAsync שולח INSERT UPDATE או DELETE. במצב מקומי התאריכים נשמרים כ TEXT של SQLite וב MySQL הם מקבלים טיפוסים מתאימים דרך migration נפרדת.")
b.h2("מפת היחסים")
b.code("""
AspNetUsers 1 ---- * Bookings * ---- 1 Resources
      |                            |      |
      |                            |      +---- * Units
      |                            |      +---- * Sessions
      +---- * RefreshSessions      |
      +---- * Notifications        +---- optional UnitId or SessionId

Bookings -> Outbox -> Kafka -> Inbox -> Notifications
Admin changes -> Audit
Retries -> IdempotencyEntries
""")
sql_tables = [
    ["Resources", "קטגוריית חדר פעילות או תחבורה", "Id Kind Name Slug Description ImageUrl GalleryImagesJson BasePriceCents MaxGuests IsActive OfferVersion FeaturesJson", "1 לרבים אל Units ו Sessions וגם Bookings"],
    ["Units", "יחידה פיזית כגון Suite 101 או Vehicle 1", "Id ResourceId Name IsActive", "ResourceId אל Resources RESTRICT"],
    ["Sessions", "מועד פעילות עם קיבולת ומחיר", "Id ResourceId StartUtc EndUtc Capacity PriceCents Version", "ResourceId אל Resources RESTRICT"],
    ["Bookings", "הזמנה ושמירת צילום התנאים", "Id BookingNumber UserId ResourceId UnitId SessionId Kind Status זמנים Guests TotalCents PriceLinesJson Version", "אל User Resource ובאופן אופציונלי Unit או Session"],
    ["IdempotencyEntries", "שמירת תגובה לפי מפתח בקשה", "UserId Operation Key RequestHash ResponseJson CreatedAtUtc", "Unique על UserId Operation Key"],
    ["Outbox", "אירועים שממתינים לפרסום ולעיבוד", "Type UserId BookingId BookingVersion PayloadJson Status Attempts זמני retry ואינדוקס", "אין FK מכוון כדי לשמור אירוע גם לאורך זמן"],
    ["Inbox", "מזהי אירועים שכבר עובדו", "Id ProcessedAtUtc", "Primary key על Event Id"],
    ["Notifications", "הודעה למשתמש ומצב מסירה", "EventId UserId Subject Message IsRead DeliveredAtUtc Lease", "EventId unique"],
    ["Audit", "מי שינה מה ומתי", "Action EntityType EntityId OldValue NewValue AdminEmail CorrelationId", "Snapshot JSON בלי FK"],
    ["RefreshSessions", "Sessions ארוכים להתחברות", "UserId TokenHash ExpiresAtUtc RevokedAtUtc", "TokenHash unique"],
    ["AspNetUsers", "משתמשי Identity ופרטי פרופיל", "Id Email PasswordHash SecurityStamp Lockout FirstName LastName Phone", "מקור קשר ל Bookings ול Identity tables"],
    ["AspNetRoles", "Roles כגון Admin ו Guest", "Id Name NormalizedName ConcurrencyStamp", "רבים לרבים עם Users"],
    ["AspNetUserRoles", "טבלת חיבור Users Roles", "UserId RoleId", "FK CASCADE לשתי הטבלאות"],
    ["AspNetUserClaims", "Claims קבועים של משתמש", "Id UserId ClaimType ClaimValue", "FK CASCADE ל Users"],
    ["AspNetRoleClaims", "Claims של Role", "Id RoleId ClaimType ClaimValue", "FK CASCADE ל Roles"],
    ["AspNetUserLogins", "כניסות מספק חיצוני אם יתווספו", "LoginProvider ProviderKey UserId", "FK CASCADE ל Users"],
    ["AspNetUserTokens", "Tokens של Identity providers", "UserId LoginProvider Name Value", "FK CASCADE ל Users"],
    ["EFMigrationsHistory", "אילו migrations כבר הוחלו", "MigrationId ProductVersion", "מנוהלת על ידי EF Core"],
]
b.table(["טבלה", "מטרה", "עמודות מרכזיות", "יחסים ואילוצים"], sql_tables, font_size=7.5)
b.h2("למה כסף נשמר ב Cents")
b.p("BasePriceCents ו TotalCents הם long ולא double. דולר 695 נשמר כ 69500. כך אין שגיאת floating point. ה API מחלק ב 100 ומחזיר decimal. BookingRules.ToCents בודקת טווח ושתי ספרות לכל היותר.")
b.h2("ירושה בטבלת Bookings")
b.p("RoomBooking ActivityBooking ו TransportationBooking יורשות מ Booking. EF משתמש ב Table Per Hierarchy: כל הסוגים נמצאים בטבלה אחת ועמודת Kind היא discriminator. היתרון הוא שאפשר לשאול את כל ההזמנות יחד. השדות שאינם מתאימים לסוג מסוים נשארים null.")
b.h2("אינדקסים")
b.bullets([
    "BookingNumber ייחודי כדי שמספר הזמנה לא יחזור.",
    "ResourceId Status StartUtc InventoryEndUtc מאיץ בדיקת חפיפות מלאי.",
    "UserId CreatedAtUtc מאיץ את מסך ההזמנות שלי.",
    "UserId Operation Key ייחודי מונע אותה פעולת idempotency פעמיים.",
    "Outbox Status NextAttemptUtc מאיץ איתור האירוע הבא לעיבוד.",
    "Slug TokenHash ו Notification EventId מקבלים Unique index."
])

b.h1("13 טרנזקציות מקביליות ו Idempotency")
b.h2("DatabaseWrite")
b.p("DatabaseWrite פותחת טרנזקציה ומקבלת נעילה לפני קריאת מלאי משתנה. ב SQLite נעשה BEGIN IMMEDIATE ברמת Serializable. ב MySQL נפתחת ReadCommitted ואז SELECT FOR UPDATE על המשתמש ועל המשאב. סדר הנעילה קבוע: קודם User ואז Resource כדי לצמצם deadlocks.")
b.p("HoldAsync פותחת את הטרנזקציה לפני בדיקת availability. לכן שתי בקשות מקבילות לא יכולות לראות בו זמנית את אותה יחידה כפנויה ואז שתיהן לשמור אותה. בסיום SaveChangesAsync כותב Booking Outbox ו IdempotencyEntry ו CommitAsync מאשר את כולם יחד.")
b.h2("Idempotency")
b.p("לקוח עלול ללחוץ פעמיים או לא לקבל תשובה למרות שהשרת הצליח. לכן POST bookings שולח Idempotency-Key. השרת מחשב SHA256 של payload ומחפש User Operation Key. אם נמצא אותו hash מוחזרת התגובה המקורית. אם אותו key הגיע עם payload אחר מוחזר 409 idempotency_conflict.")
b.h2("Optimistic concurrency")
b.p("OfferVersion של Resource Version של Session ו Version של Booking מוגדרים ConcurrencyToken. ה Frontend שולח ExpectedVersion שקרא קודם. שירות הניהול משווה ומעלה Version. אם מנהל אחר שמר בינתיים מתקבל 409 והמסך צריך לטעון מחדש. זה מונע lost update.")

b.h1("14 Outbox Inbox Kafka Elasticsearch ו Kibana")
b.h2("הבעיה שה Outbox פותר")
b.p("נניח שהזמנה נשמרה ב MySQL ואז האפליקציה ניסתה לשלוח הודעה ל Kafka ונפלה. ההזמנה קיימת אך אין אירוע. אם שולחים קודם ל Kafka ואז השמירה נכשלת נוצר אירוע להזמנה שאינה קיימת. Outbox פותר את הפער: Booking וה OutboxEvent נשמרים באותה טרנזקציית SQL. לאחר Commit מובטח שהאירוע קיים לעיבוד עתידי.")
b.h2("זרימת האירוע")
b.bullets([
    "BookingService.AddEvent מוסיפה OutboxEvent אך אינה פונה לרשת.",
    "HotelWorker.PublishOneAsync נועל אירוע Pending מסמן Processing ונותן lease של דקה.",
    "לאחר Commit הוא מפרסם EventEnvelope ל topic vesper.bookings ב Kafka.",
    "Producer מוגדר EnableIdempotence ו Acks All. מפתח ההודעה הוא BookingId כדי לשמור סדר בתוך partition.",
    "Consumer בקבוצה vesper-notifications-v1 קורא את ההודעה בלי Auto Commit.",
    "ConsumeInTransactionAsync בודק Inbox. אם Id כבר קיים הוא אינו יוצר Notification נוספת.",
    "אם האירוע חדש נוצרים Notification ו InboxEvent באותה טרנזקציה ורק אז Kafka offset מקבל Commit.",
    "DeliverOneAsync שומר הודעת JSON בתיקיית notifications כתחליף לספק דואר אמיתי.",
    "IndexOneAsync שולח metadata של האירוע לאינדקס vesper-events ב Elasticsearch."
])
b.h2("Retries ו Leases")
b.p("Outbox מחזיק Attempts NextAttemptUtc LastError ו Status. כשל מחזיר את האירוע ל Pending עם exponential backoff עד 300 שניות. אחרי עשרה ניסיונות הוא Failed. Processing עם NextAttemptUtc בעתיד הוא lease מתמשך: אם Worker נופל אירוע חוזר להיות מועמד לאחר דקה. אינדוקס ל Elasticsearch מחזיק מונה וזמן retry נפרדים כך שכשל בחיפוש אינו חוסם התראה.")
b.h2("Kafka")
b.p("Kafka אינו מסד ההזמנות ואינו נדרש לקריאת קטלוג. הוא מפריד בין יצירת האירוע לעיבודו. Producer כותב topic ו Consumer group קוראת. Offset הוא המיקום של הקבוצה ב topic. בקוד ה offset מאושר רק לאחר כתיבת Inbox ו Notification. אם התהליך נופל לפני commit Kafka תמסור שוב וה Inbox תסנן כפילות.")
b.h2("Elasticsearch ו Kibana")
b.p("Elasticsearch מקבל מסמך קצר לכל אירוע שאינו PasswordReset: Id Type BookingId BookingVersion ו CreatedAtUtc. המטרה היא חיפוש ותצפית ולא שחזור מצב ההזמנה. Kibana מתחברת ל Elasticsearch ומאפשרת ליצור Data View עבור vesper-events לחפש BookingId לספור סוגי אירועים ולראות כשלים לאורך זמן. מחיקת אינדקס Elastic אינה מוחקת הזמנה כי MySQL נשאר מקור האמת.")
b.h2("מצב מקומי ללא Kafka")
b.p("LocalHotelWorker משתמש באותן טבלאות Outbox Inbox Notifications אך מעבד אותן ישירות בתוך ה API. הדבר מאפשר ללמוד ולהריץ ללא Docker. מצב Docker מכבה Workers RunInApi ומפעיל את Hotel.Worker הנפרד עם Kafka ו Elasticsearch.")

b.h1("15 ה API החיצוני של מזג האוויר")
b.p("WeatherService מקבל HttpClient מ Dependency Injection. BaseAddress מגיע מ Weather BaseUrl וה timeout הוא שמונה שניות. GetLasVegasAsync קוראת Open Meteo Forecast API עם קואורדינטות לאס וגאס יחידות Fahrenheit ו mph אזור זמן America Los Angeles ותחזית לארבעה ימים.")
b.h2("מיפוי JSON")
b.p("OpenMeteoResponse OpenMeteoCurrent ו OpenMeteoDaily מייצגות בדיוק את שמות השדות של הספק בעזרת JsonPropertyName. השירות ממיר אותן ל WeatherResponse פנימי. כך ה Frontend אינו תלוי ישירות במבנה הספק ואפשר להחליף ספק בעתיד בתוך WeatherService.")
b.h2("Cache וטיפול בכשל")
b.p("התגובה נשמרת ב IMemoryCache למשך 15 דקות תחת המפתח las-vegas-weather. הדבר מצמצם latency ותעבורה לספק. WeatherController תופס HttpRequestException או timeout ומחזיר 503 weather_unavailable. ב Frontend קיימת גם fallback ישירה ל Open Meteo עבור סביבת Windows מקומית שבה Schannel עשוי לחסום את קריאת השרת. ב Docker Linux הזרימה הראשית היא דרך ה API.")
b.h2("עקרונות לשימוש ב API חיצוני")
b.bullets([
    "מגדירים timeout ולא מחכים לנצח.",
    "לא חושפים את מבנה הספק לכל המערכת אלא ממפים ל DTO פנימי.",
    "משתמשים ב cache כאשר הנתון אינו חייב להיות עדכני בכל שנייה.",
    "מטפלים ב 503 ומציגים UI תקין גם אם הספק נפל.",
    "מפתחות API אם יידרשו בעתיד יישמרו ב Environment variables ולא בקוד."
])

b.h1("16 Migrations ו Seed")
b.h2("מהי Migration")
b.p("Migration היא מחלקת C# עם Up ו Down. Up משנה את schema קדימה. Down מתאר חזרה לאחור. EF Core משווה את ModelSnapshot למודל החדש ומייצר פעולות כמו CreateTable AddColumn CreateIndex. הטבלה EFMigrationsHistory אומרת אילו migrations כבר רצו ולכן כל migration מוחלת פעם אחת.")
b.h2("למה יש שתי סדרות")
b.p("SQLite ו MySQL משתמשים בטיפוסים ובתחביר שונים. migrations של SQLite נמצאות ב Hotel.Infrastructure Migrations. migrations של MySQL נמצאות ב Hotel.Migrations.MySql וה DbContext בוחר MigrationsAssembly לפי Database Provider. כך Docker אינו מנסה להריץ migration שנוצרה במיוחד ל SQLite.")
b.table(["Migration", "Provider", "מה היא עושה"], [
    ["20261008090820 InitialCreate", "SQLite", "יוצרת Identity וכל טבלאות המלון יחסים ואינדקסים"],
    ["20261008092006 InitialMySql", "MySQL", "אותו מודל עם טיפוסי MySQL ומגבלות מפתחות"],
    ["20261011061049 AddResourceGalleries", "SQLite", "מוסיפה GalleryImagesJson ל Resources"],
    ["20261011061109 AddResourceGalleriesMySql", "MySQL", "מוסיפה nullable ממלאת [] ואז משנה ל non null"],
])
b.h2("DatabaseSeeder")
b.p("לאחר Migrate Program מפעיל InitializeAsync. ה Seeder יוצר Roles Admin ו Guest יוצר Admin רק אם אינו קיים יוצר שורת system לנעילות וזורע חדרים פעילויות רכבים יחידות ומועדים במסד ריק. AddDemoRoomGalleriesAsync מוסיפה תמונות לחדרי הדוגמה ומנקה ערכי גלריה ישנים ריקים. Seed נכתב להיות idempotent כדי שלא ייצור קטלוג כפול בכל הפעלה.")
b.h2("איך יוצרים Migration חדשה")
b.code("""
dotnet ef migrations add NameOfChange \
  --project src/Hotel.Infrastructure \
  --startup-project src/Hotel.Api

dotnet ef migrations add NameOfChangeMySql \
  --project src/Hotel.Migrations.MySql \
  --startup-project src/Hotel.Api
""")
b.p("לפני שמירת migration יש לקרוא את Up ו Down במיוחד בשינוי עמודה קיימת. שינוי non null דורש אסטרטגיה לנתונים קיימים כמו שנעשה בגלריות: הוספה nullable עדכון ערכים ואז מעבר ל non null.")

b.h1("17 הזרימות המלאות")
b.h2("זרימת טעינת דף הבית")
b.bullets([
    "הדפדפן טוען index.html ואת bundle של Vite.", "main.tsx יוצר Router ו AuthProvider.",
    "AuthProvider מנסה POST auth refresh כדי לשחזר session.", "HomePage קוראת GET catalog kind Room ו WeatherWidget קוראת GET weather.",
    "CatalogController מפעיל CatalogService שמבצע LINQ על Resources Units Sessions ו ActiveBookings.",
    "ה JSON חוזר ל useApi State מתעדכן ו React מצייר כרטיסי חדרים ותמונות."
], numbered=True)
b.h2("זרימת יצירת הזמנה")
b.bullets([
    "BookingPanel קורא availability עם resource dates guests session.",
    "BookingRules ממירה זמנים ובודקת אורחים והקצאה מחשבת מחיר ומלאי.",
    "בעת Reserve ה Frontend יוצר Idempotency Key ושולח HoldRequest עם OfferVersion.",
    "BookingsController קורא UserId מה JWT ומפעיל HoldAsync.",
    "DatabaseWrite נועלת User ו Resource בתוך transaction.",
    "ReplayAsync בודקת אם הבקשה כבר טופלה.",
    "השירות קורא מחדש Resource ובודק סוג Version תאריכים אורחים ו route.",
    "AllocateAsync בוחרת Unit פנויה או בודקת מושבי Session.",
    "נוצר Booking במצב Pending לעשר דקות ונשמרים המחיר והתנאים.",
    "נוצרים גם OutboxEvent ו IdempotencyEntry באותה טרנזקציה.",
    "לאחר Commit מוחזרת BookingResponse והדפדפן עובר לעמוד ההזמנה."
], numbered=True)
b.h2("זרימת אישור")
b.p("POST bookings id confirm נכנסת ל TransitionAsync. השירות נועל את המשאב בודק שה Hold לא פג ושמצבו Pending משנה ל Confirmed מבטל ExpiresAtUtc מעלה Version ומוסיף BookingConfirmed ל Outbox. המחיר אינו מחושב מחדש ולכן נשמר המחיר שננעל.")
b.h2("זרימת ביטול")
b.p("TransitionAsync בודקת בעלות או Admin מצב ההזמנה זמן הסיום ו CancellationDeadlineUtc. משתמש רגיל נחסם אחרי הדדליין. Admin יכול לעקוף אך חייב reason ונוצרת AuditEntry. הסטטוס עובר ל Cancelled ונוצר אירוע.")
b.h2("זרימת פקיעה והשלמה")
b.p("ה Worker מחפש Resources שבהם Pending עבר את ExpiresAtUtc או Confirmed עבר את EndUtc. SweepResourceAsync נועלת כל Resource בנפרד משנה Pending ל Expired ו Confirmed ל Completed ומייצרת אירוע מתאים. כך מלאי חוזר לזמינות גם בלי בקשת משתמש.")
b.h2("זרימת שינוי מלאי באדמין")
b.p("AdminResources טוען גם משאבים לא פעילים. ResourceEditor שולח ExpectedVersion. AdminController מוגן ב Role Admin. AdminService נועלת Resource מאמתת נתונים משווה Version שומרת ומוסיפה Audit באותה טרנזקציה. לאחר הצלחה ה Frontend מרענן את הרשימה.")
b.h2("זרימת Forgot Password")
b.p("המשתמש שולח Email. AuthService מחזירה תמיד תשובה כללית. אם החשבון קיים Identity מייצר token לשעה וקישור. האירוע PasswordReset נשמר ב Outbox ה Worker הופך אותו ל Notification ו DeliverOneAsync שומר קובץ JSON בתיקיית notifications. בסביבת אמת מחליפים שלב זה בספק Email.")

b.h1("18 Docker בדיקות ותחזוקה")
b.h2("שירותי Docker Compose")
b.table(["Service", "תפקיד", "פורט מקומי"], [
    ["web", "nginx שמגיש React ומעביר api ל API", "8080"],
    ["api", "ASP.NET Core Controllers ו services", "5080"],
    ["worker", "Kafka notifications lifecycle ו Elastic indexing", "פנימי"],
    ["mysql", "מסד מקור האמת", "3307"],
    ["kafka", "Broker אירועים במצב KRaft", "9094"],
    ["elasticsearch", "אינדקס אירועים", "9200"],
    ["kibana", "ממשק תצפית", "5601"],
])
b.code("""
powershell -ExecutionPolicy Bypass -File scripts/initialize-env.ps1
docker compose --env-file .env up --build

# עצירה תוך שמירת volumes
powershell -ExecutionPolicy Bypass -File scripts/stop-docker.ps1
""")
b.h2("Environment variables חשובים")
b.table(["משתנה", "שימוש"], [
    ["MYSQL_PASSWORD ו MYSQL_ROOT_PASSWORD", "גישה למסד Docker"],
    ["JWT_KEY", "חתימת Access tokens לפחות 32 bytes"],
    ["ADMIN_EMAIL ו ADMIN_PASSWORD", "Seed של חשבון מנהל"],
    ["KAFKA_CLUSTER_ID", "זהות broker במצב KRaft"],
    ["Database Provider", "Sqlite או MySql"],
    ["Messaging Mode", "Local או Kafka"],
    ["Weather BaseUrl", "כתובת Open Meteo"],
    ["Elasticsearch Url", "יעד אינדוקס אירועים"],
])
b.h2("בדיקות")
b.p("tests api-smoke.mjs היא בדיקת אינטגרציה. היא נכנסת כמנהל יוצרת משתמשים ומשאבים זמניים ומאמתת הרשאות Hold idempotency תחרות על חדר אחרון קיבולת פעילות buffer לרכב נעילת מחיר Dashboard ו Audit. היא גם בודקת שלחדרי ההדגמה יש תמונות שער שונות וגלריות. בסוף היא משביתה את משאבי הבדיקה.")
b.h2("סדר מומלץ לאיתור תקלה")
b.bullets([
    "בדוק api health ו docker compose ps.",
    "בדוק status code ו response JSON ב Network של הדפדפן.",
    "השתמש ב traceId כדי לחפש את אותה בקשה בלוג ה API.",
    "בדוק אם Validation החזירה 400 אם Version החזיר 409 או אם Authentication החזירה 401 או 403.",
    "בתקלת אירוע בדוק Admin System ואז Outbox Status Attempts LastError.",
    "בדוק Kafka consumer רק לאחר שווידאת שה Outbox נכתבה.",
    "בדוק Elasticsearch בנפרד משום שכשל בו אינו אמור לשבור הזמנה או הודעה."
])

# Full backend function catalogue.
b.h1("נספח א כל פונקציות C#")
b.p("הרשימה כוללת את כל המתודות וה constructors בעלי משמעות בקוד שנכתב לפרויקט. Accessors אוטומטיים של properties ומתודות generated של migrations אינם מפורטים כשורות נפרדות. מתודות Up ו Down של migrations מוסברות בפרק 16.")
backend_functions = [
    ["BookingRules.ToCents", "decimal price", "long cents", "מאמתת מחיר וממירה דולרים לסנטים ללא floating point"],
    ["BookingRules.Period", "kind start end now", "BookingPeriod", "מפרשת תאריך חדר או timestamps מחשבת Quantity UTC ו InventoryEndUtc"],
    ["BookingRules.ValidateGuests", "guests maxGuests", "void", "זורקת שגיאה אם הכמות מחוץ לטווח"],
    ["BookingRules.Overlaps", "שני טווחי זמן", "bool", "בודקת חפיפה לפי start less than other end"],
    ["DatabaseWrite.BeginAsync", "DbContext resourceId userId token", "DatabaseWrite", "פותחת transaction ונועלת User ו Resource בסדר קבוע"],
    ["DatabaseWrite.CommitAsync", "CancellationToken", "Task", "מבצעת Commit"],
    ["DatabaseWrite.DisposeAsync", "ללא", "ValueTask", "משחררת transaction וחיבור SQLite"],
    ["CatalogService.ListAsync", "kind includeInactive", "List CatalogItem", "טוענת קטלוג יחידות Sessions ומושבים שמורים"],
    ["CatalogService.GetAsync", "id admin", "CatalogItem", "טוענת משאב בודד ומסתירה inactive מאורח"],
    ["CatalogService.Map", "entity counts", "CatalogItem", "ממפה Entity ל DTO"],
    ["CatalogService.ReadStringList", "json", "List string", "קוראת JSON בבטחה ומחזירה רשימה ריקה לנתון ישן פגום"],
    ["CatalogService.Session", "ActivitySession reserved", "SessionResponse", "ממפה Session ומוסיף ReservedSeats"],
    ["BookingService.ActiveBookings", "now", "IQueryable Booking", "מגדירה Confirmed או Pending שלא פג כצורכי מלאי"],
    ["BookingService.AvailabilityAsync", "resource dates guests session", "AvailabilityResponse", "מאמתת ומחשבת מלאי מחיר וגרסת הצעה בלי לשמור"],
    ["BookingService.HoldAsync", "user request key", "BookingResponse", "נועלת מלאי יוצרת Pending Booking Outbox ו idempotency"],
    ["BookingService.TransitionAsync", "id user admin action reason key audit data", "BookingResponse", "מאשרת או מבטלת ומאמתת מעבר מצב"],
    ["BookingService.GetAsync", "id user admin", "BookingResponse", "מחזירה הזמנה רק לבעלים או מנהל"],
    ["BookingService.ListAsync", "filters paging", "PagedBookings", "מסננת ומעמדת הזמנות"],
    ["BookingService.SweepResourceAsync", "resourceId", "Task", "מסמנת Holds שפגו והזמנות שהסתיימו ומייצרת אירועים"],
    ["BookingService.AllocateAsync", "resource period guests sessionId", "Allocation", "מוצאת Unit פנויה או בודקת קיבולת Session"],
    ["BookingService.ReplayAsync", "user operation key hash", "BookingResponse nullable", "מחזירה תגובה קודמת או מזהה שימוש שונה במפתח"],
    ["BookingService.SaveReplay", "user operation key hash response", "void", "מוסיף IdempotencyEntry"],
    ["BookingService.ValidateKey", "key", "void", "דורשת מפתח באורך תקין"],
    ["BookingService.Hash", "string", "SHA256 hex", "מייצרת hash ל payload או token"],
    ["BookingService.AddEvent", "booking type", "void", "מוסיפה OutboxEvent עם NotificationMessage"],
    ["AdminService.SaveResourceAsync", "id request email correlation", "CatalogItem", "יוצרת או מעדכנת משאב עם validation version ו audit"],
    ["AdminService.SaveUnitAsync", "id resourceId request audit data", "ResourceUnit", "יוצרת מעדכנת או משביתה יחידה פיזית"],
    ["AdminService.SaveSessionAsync", "id resourceId request audit data", "SessionResponse", "מנהלת מועד פעילות קיבולת מחיר ו version"],
    ["AdminService.DashboardAsync", "from to", "object", "מחשב ערך הזמנות תפוסה ניצול סדרות יומיות והזמנות קרובות"],
    ["AdminService.Audit", "action entity snapshots admin correlation", "void", "מוסיף AuditEntry"],
    ["AuthService.RegisterAsync", "RegisterRequest", "AuthSession", "יוצרת Guest ומנפיקה tokens"],
    ["AuthService.LoginAsync", "LoginRequest", "AuthSession", "בודקת סיסמה ונעילה ומנפיקה tokens"],
    ["AuthService.RefreshAsync", "refresh token", "AuthSession", "מאמתת hash מבטלת token ישן ומבצעת rotation"],
    ["AuthService.LogoutAsync", "refresh token", "Task", "מבטלת RefreshSession"],
    ["AuthService.MeAsync", "user id", "UserResponse", "מחזירה משתמש נוכחי"],
    ["AuthService.ProfileAsync", "id ProfileRequest", "UserResponse", "מעדכנת שם וטלפון"],
    ["AuthService.ForgotAsync", "email", "Task", "מייצרת reset link דרך Outbox בלי לחשוף קיום משתמש"],
    ["AuthService.ResetAsync", "ResetPasswordRequest", "Task", "מאפסת סיסמה ומבטלת sessions"],
    ["AuthService.ListUsersAsync", "page pageSize search", "paged object", "מחזירה רשימת משתמשים לאדמין"],
    ["AuthService.IssueAsync", "HotelUser", "AuthSession", "יוצרת JWT ו Refresh token hash"],
    ["AuthService.MapAsync", "HotelUser", "UserResponse", "ממפה משתמש ומחשב Role"],
    ["WeatherService.GetLasVegasAsync", "CancellationToken", "WeatherResponse", "קוראת Open Meteo או cache וממפה תגובה"],
    ["WeatherService.WeatherCondition", "WMO code", "string", "ממירה קוד מזג אוויר לטקסט"],
    ["DatabaseSeeder.InitializeAsync", "IServiceProvider", "Task", "מריצה migrations יוצרת Roles Admin ונתוני הדגמה"],
    ["DatabaseSeeder.Resource", "שדות משאב", "CatalogResource", "Factory משותף ל seed"],
    ["DatabaseSeeder.AddRoom", "פרטי חדר", "void", "זורעת חדר גלריה ויחידות"],
    ["DatabaseSeeder.AddDemoRoomGalleriesAsync", "DbContext", "Task", "משדרגת נתוני גלריה קיימים"],
    ["DatabaseSeeder.AddActivity", "פרטי פעילות", "void", "זורעת פעילות ו Sessions עתידיים"],
    ["DatabaseSeeder.AddTransport", "פרטי רכב", "void", "זורעת שירות תחבורה ויחידות"],
    ["ServiceRegistration.AddHotel", "services config", "IServiceCollection", "רושמת DB Identity JWT services cache ו HttpClient"],
    ["JsonDefaults.Serialize", "object", "string", "JSON camelCase עם enums כמחרוזות"],
    ["BookingResponse.From", "Booking", "BookingResponse", "ממפה Entity ל DTO וממיר cents"],
    ["ApiErrors.InvokeAsync", "HttpContext", "Task", "Middleware לשגיאות headers logging ו correlation"],
    ["HotelWorker.ExecuteAsync", "stoppingToken", "Task", "לולאת sweep publish consume deliver index"],
    ["HotelWorker.SweepAsync", "token", "Task", "מאתרת resources עם הזמנות לפקיעה או השלמה"],
    ["HotelWorker.PublishOneAsync", "producer topic token", "Task", "claims Outbox מפרסם Kafka ומעדכן retries"],
    ["HotelWorker.ConsumeAsync", "envelope token", "Task", "פותחת transaction לעיבוד consumer"],
    ["HotelWorker.ConsumeInTransactionAsync", "db envelope token", "Task", "Inbox dedupe ויצירת Notification"],
    ["HotelWorker.DeliverOneAsync", "token", "Task", "claims notification וכותב קובץ מסירה"],
    ["HotelWorker.IndexOneAsync", "token", "Task", "מאנדקס אירוע ב Elasticsearch עם retry עצמאי"],
    ["HotelWorker.Dispose", "ללא", "void", "משחררת HttpClient"],
    ["LocalHotelWorker.ExecuteAsync", "token", "Task", "לולאת רקע מקומית ללא Kafka"],
    ["LocalHotelWorker.SweepBookingsAsync", "token", "Task", "מפעילה sweep לפי resource"],
    ["LocalHotelWorker.ProcessOutboxAsync", "token", "Task", "Outbox אל Inbox ו Notification באותה DB"],
    ["LocalHotelWorker.DeliverNotificationAsync", "token", "Task", "כותבת הודעה מקומית ומסמנת Delivered"],
    ["AdminController endpoints", "HTTP params", "Action results", "Dashboard Bookings Resources Units Sessions Users Audit SystemStatus"],
    ["AuthController Register Login Refresh Logout", "request או cookie", "AuthResponse", "עוטפת AuthService ומנהלת Refresh cookie"],
    ["AuthController Me Profile Forgot Reset", "claims ו DTOs", "responses", "פרופיל ואיפוס סיסמה"],
    ["AuthController.CookieOptions", "ללא", "CookieOptions", "HttpOnly SameSite Strict path auth ושבעה ימים"],
    ["AuthController.SetCookie", "token", "void", "כותבת vesper.refresh"],
    ["BookingsController Hold Mine Get Confirm Cancel", "HTTP ו claims", "Booking DTOs", "ממפה API לתהליכי BookingService"],
    ["CatalogController List Get Availability", "query route", "catalog DTOs", "API ציבורי לקטלוג וזמינות"],
    ["NotificationsController.List", "User claim", "notifications", "מחזירה 100 הודעות אחרונות של המשתמש"],
    ["NotificationsController.Read", "notification id", "204", "נועלת ומסמנת הודעה של המשתמש כנקראה"],
    ["WeatherController.Get", "token", "WeatherResponse או 503", "עוטפת WeatherService בטיפול כשל ספק"],
    ["Program health endpoint", "DbContext config", "health JSON", "בודק CanConnect ומדווח provider ו messaging mode"],
]
b.table(["פונקציה", "קלט", "פלט", "מה היא עושה"], backend_functions, font_size=7.2)

b.h1("נספח ב כל פונקציות Frontend")
b.p("הרשימה כוללת Components Hooks ופונקציות עזר בעלות שם. callbacks קצרים בתוך JSX כגון onClick מתוארים כחלק מה Component שמכיל אותם.")
frontend_functions = [
    ["setAccessToken", "api.ts", "שומר או מנקה Access token בזיכרון"],
    ["api", "api.ts", "שולחת fetch מוסיפה headers מנסה refresh ומתרגמת שגיאות"],
    ["refreshToken", "api.ts", "מבקשת Access token חדש בעזרת cookie"],
    ["send", "api.ts", "עטיפה ל POST PUT עם JSON body"],
    ["money", "api.ts", "מפרמט סכום לפי מטבע"],
    ["hotelDate", "api.ts", "מציג תאריך באזור הזמן של לאס וגאס"],
    ["errorMessage", "api.ts", "מחלץ הודעת Error בטוחה"],
    ["upcomingDate", "api.ts", "מחזיר YYYY-MM-DD בעוד מספר ימים"],
    ["vegasToUtc", "api.ts", "ממיר datetime-local של לאס וגאס ל UTC כולל DST"],
    ["utcToVegasInput", "api.ts", "ממיר UTC לערך datetime-local"],
    ["AuthProvider", "auth.tsx", "מנהל user loading bootstrap refresh login acceptance ו logout"],
    ["accept", "בתוך AuthProvider", "שומר token ו user לאחר login"],
    ["logout", "בתוך AuthProvider", "קורא logout ומנקה state"],
    ["useAuth", "auth.tsx", "מחזיר AuthContext"],
    ["useApi", "components.tsx", "Hook לטעינת API עם data error loading refresh"],
    ["Loading", "components.tsx", "מצב טעינה נגיש"],
    ["ErrorBox", "components.tsx", "שגיאה וכפתור retry"],
    ["SuccessBox", "components.tsx", "הודעת הצלחה"],
    ["Empty", "components.tsx", "מצב רשימה ריקה"],
    ["Status", "components.tsx", "badge לפי ערך מצב"],
    ["PageTitle", "components.tsx", "כותרת עמוד אחידה"],
    ["Modal", "components.tsx", "dialog עם Escape נעילת גלילה וסגירה ברקע"],
    ["Pagination", "components.tsx", "טווח תוצאות Previous Next"],
    ["ResourceCard", "components.tsx", "כרטיס קטלוג עם תמונה מחיר ותכונות"],
    ["ScrollToTop", "App.tsx", "גולל לראש לאחר שינוי route"],
    ["Guard", "App.tsx", "מגן routes לפי login ו Admin role"],
    ["PublicLayout", "App.tsx", "Header ניווט מובייל Outlet ו Footer"],
    ["Footer", "App.tsx", "קישורי אתר וקישור Hotel management"],
    ["AboutPage", "App.tsx", "עמוד תוכן ומדיניות"],
    ["NotFound", "App.tsx", "מסך 404"],
    ["App", "App.tsx", "מגדיר את כל Routes"],
    ["StaySearch", "PublicPages.tsx", "שומר dates guests ומנווט ל rooms עם query"],
    ["WeatherWidget", "PublicPages.tsx", "טוען מזג אוויר API עם fallback ומציג forecast"],
    ["HomePage", "PublicPages.tsx", "מרכיב את דף הבית"],
    ["CatalogPage", "PublicPages.tsx", "טוען מסנן וממיין Resources"],
    ["DetailPage", "PublicPages.tsx", "טוען Resource ומרכיב גלריה ו BookingPanel"],
    ["RoomGallery", "PublicPages.tsx", "גלריית חדר lightbox keyboard וניווט"],
    ["BookingPanel", "PublicPages.tsx", "זמינות quote יצירת Hold וניווט"],
    ["AuthShell", "AuthPages.tsx", "מעטפת חזותית למסכי auth"],
    ["PasswordInput", "AuthPages.tsx", "input עם show hide"],
    ["AuthPage", "AuthPages.tsx", "Login או Register ושמירת AuthContext"],
    ["AdminLoginPage", "AuthPages.tsx", "Login ייעודי ובדיקת Admin role"],
    ["ForgotPasswordPage", "AuthPages.tsx", "שליחת forgot password"],
    ["ResetPasswordPage", "AuthPages.tsx", "קורא query ושולח password חדש"],
    ["AccountLayout", "AccountPages.tsx", "ניווט account ו Outlet"],
    ["BookingCard", "AccountPages.tsx", "סיכום הזמנה"],
    ["BookingsPage", "AccountPages.tsx", "רשימת הזמנות משתמש וסינון"],
    ["ProfilePage", "AccountPages.tsx", "עריכת פרופיל"],
    ["NotificationsPage", "AccountPages.tsx", "רשימת התראות וסימון Read"],
    ["CancelBookingModal", "AccountPages.tsx", "ביטול והצגת שגיאה"],
    ["HoldTimer", "AccountPages.tsx", "ספירה לאחור ועדכון Expired UI"],
    ["BookingPage", "AccountPages.tsx", "פרטי Hold Confirm Cancel"],
    ["AdminLayout", "AdminPages.tsx", "Sidebar Topbar Outlet ופרטי מנהל"],
    ["AdminTitle", "AdminPages.tsx", "כותרת מסך ניהול"],
    ["Metric", "AdminPages.tsx", "כרטיס מדד"],
    ["percent", "AdminPages.tsx", "מפרמט יחס או אחוז"],
    ["AdminDashboard", "AdminPages.tsx", "מסנני זמן metrics chart ו upcoming"],
    ["BookingTable", "AdminPages.tsx", "טבלת הזמנות משותפת"],
    ["AdminBookings", "AdminPages.tsx", "חיפוש סינון צפייה וביטול מנהלי"],
    ["AdminResources", "AdminPages.tsx", "רשימת משאבים ופתיחת עורכים"],
    ["ResourceEditor save", "AdminPages.tsx", "יוצר או מעדכן משאב גלריה features version"],
    ["InventoryEditor", "AdminPages.tsx", "בוחר UnitManager או SessionManager"],
    ["UnitManager save toggle", "AdminPages.tsx", "הוספה שינוי שם הפעלה והשבתה"],
    ["SessionManager edit save", "AdminPages.tsx", "עריכת מועד המרת זמנים version ו capacity"],
    ["AdminUsers", "AdminPages.tsx", "חיפוש ומשתמשים בעמודים"],
    ["AdminAudit", "AdminPages.tsx", "Audit list ו before after modal"],
    ["formatValue", "AdminPages.tsx", "Pretty print ל JSON snapshot"],
    ["AdminSystem", "AdminPages.tsx", "health outbox notifications ו recent events"],
]
b.table(["פונקציה או Component", "קובץ", "אחריות"], frontend_functions, font_size=7.5)

b.h1("נספח ג כל האובייקטים")
objects = [
    ["ResourceKind", "enum", "Room Activity Transportation"],
    ["BookingStatus", "enum", "Pending Confirmed Cancelled Completed Expired Failed"],
    ["CatalogResource", "Entity", "קטגוריית מוצר מלונאי מחיר תיאור תמונות features ו version"],
    ["ResourceUnit", "Entity", "חדר או רכב פיזי השייך ל Resource"],
    ["ActivitySession", "Entity", "מועד פעילות קיבולת מחיר ו version"],
    ["Booking", "Entity abstract", "בסיס משותף לכל הזמנה כולל snapshot מחיר זמן והקצאה"],
    ["RoomBooking", "Entity", "Booking מסוג Room"],
    ["ActivityBooking", "Entity", "Booking מסוג Activity"],
    ["TransportationBooking", "Entity", "Booking מסוג Transportation"],
    ["BookingPeriod", "Domain value", "Start End InventoryEnd ו Quantity לאחר פירוש קלט"],
    ["HotelException", "Domain error", "שגיאה צפויה עם HTTP status ו code"],
    ["IdempotencyEntry", "Entity", "מפתח בקשה hash ותגובה שמורה"],
    ["OutboxEvent", "Entity", "אירוע מתמשך עם retry ומצב אינדוקס"],
    ["InboxEvent", "Entity", "Event Id שכבר עובד"],
    ["Notification", "Entity", "הודעה למשתמש מצב קריאה ומסירה"],
    ["AuditEntry", "Entity", "שינוי מנהלי before after וזהות מבצע"],
    ["RefreshSession", "Entity", "hash של refresh token תוקף וביטול"],
    ["HotelUser", "Identity Entity", "IdentityUser בתוספת FirstName ו LastName"],
    ["HoldRequest", "Request DTO", "קלט ליצירת Hold"],
    ["PriceLine", "DTO", "שורת תמחור quantity unit total"],
    ["BookingResponse", "Response DTO", "הזמנה בטוחה להצגה ב API"],
    ["PagedBookings", "Response DTO", "Items TotalCount Page PageSize"],
    ["SessionResponse", "Response DTO", "Session כולל ReservedSeats"],
    ["CatalogItem", "Response DTO", "Resource כולל רשימות מפוענחות וזמינות"],
    ["AvailabilityResponse", "Response DTO", "זמין כמות מחיר version וביטול"],
    ["ResourceRequest", "Request DTO", "עריכת משאב באדמין"],
    ["UnitRequest", "Request DTO", "שם ומצב יחידה"],
    ["SessionRequest", "Request DTO", "זמן קיבולת מחיר ו version"],
    ["CancelRequest", "Request DTO", "סיבת ביטול"],
    ["RegisterRequest", "Auth DTO", "Email Password שם וטלפון"],
    ["LoginRequest", "Auth DTO", "Email Password"],
    ["ProfileRequest", "Auth DTO", "שם וטלפון"],
    ["ForgotPasswordRequest", "Auth DTO", "Email"],
    ["ResetPasswordRequest", "Auth DTO", "Email Token NewPassword"],
    ["UserResponse", "Auth DTO", "פרטי משתמש ו Role ללא סיסמה"],
    ["AuthResponse", "Auth DTO", "AccessToken ו User"],
    ["AuthSession", "Internal DTO", "AuthResponse ו RefreshToken לפני כתיבת cookie"],
    ["NotificationMessage", "Event payload", "Subject ו Message"],
    ["EventEnvelope", "Kafka DTO", "Metadata ו PayloadJson של Outbox"],
    ["WeatherResponse", "Response DTO", "מזג אוויר נוכחי ו Daily"],
    ["DailyWeather", "DTO", "תאריך מינימום מקסימום condition"],
    ["OpenMeteoResponse", "External DTO", "current ו daily מהספק"],
    ["OpenMeteoCurrent", "External DTO", "שדות current במבנה הספק"],
    ["OpenMeteoDaily", "External DTO", "מערכי תחזית במבנה הספק"],
    ["Allocation", "Private object", "תוצאת הקצאת Unit או Session מחיר וכמות"],
    ["ApiError", "Frontend class", "Error עם status ו code"],
    ["ActivitySession interface", "Frontend type", "העתק typed של SessionResponse"],
    ["CatalogItem interface", "Frontend type", "העתק typed של catalog JSON"],
    ["User interface", "Frontend type", "משתמש מחובר"],
    ["AuthResponse interface", "Frontend type", "login response"],
    ["Availability interface", "Frontend type", "availability response"],
    ["Booking interface", "Frontend type", "booking response"],
    ["Paged T interface", "Frontend generic", "תוצאה ממועמדת מכל סוג"],
    ["Notice Unit Audit Dashboard SystemStatus Weather", "Frontend interfaces", "טיפוסים למסכי account admin ומזג אוויר"],
    ["AuthContextValue", "Frontend interface", "החוזה של AuthProvider לצאצאים"],
]
b.table(["אובייקט", "סוג", "מה הוא מייצג"], objects, font_size=7.6)

b.h1("נספח ד כל נתיבי ה API")
routes = [
    ["GET", "/api/health", "ציבורי", "בריאות DB ו messaging mode"],
    ["GET", "/api/catalog", "ציבורי", "קטלוג עם kind אופציונלי"],
    ["GET", "/api/catalog/{id}", "ציבורי", "פרטי משאב"],
    ["GET", "/api/availability", "ציבורי", "זמינות ומחיר"],
    ["GET", "/api/weather", "ציבורי", "מזג אוויר לאס וגאס"],
    ["POST", "/api/auth/register", "ציבורי rate limited", "הרשמה cookie ו access token"],
    ["POST", "/api/auth/login", "ציבורי rate limited", "כניסה"],
    ["POST", "/api/auth/refresh", "Refresh cookie", "token rotation"],
    ["POST", "/api/auth/logout", "Refresh cookie", "ביטול session"],
    ["GET", "/api/auth/me", "User", "משתמש נוכחי"],
    ["PUT", "/api/auth/profile", "User", "עדכון פרופיל"],
    ["POST", "/api/auth/forgot-password", "ציבורי rate limited", "בקשת איפוס"],
    ["POST", "/api/auth/reset-password", "ציבורי rate limited", "איפוס סיסמה"],
    ["POST", "/api/bookings", "User Idempotency Key", "יצירת Hold"],
    ["GET", "/api/bookings/mine", "User", "הזמנות משתמש"],
    ["GET", "/api/bookings/{id}", "Owner או Admin", "פרטי הזמנה"],
    ["POST", "/api/bookings/{id}/confirm", "Owner או Admin", "אישור"],
    ["POST", "/api/bookings/{id}/cancel", "Owner או Admin", "ביטול"],
    ["GET", "/api/notifications", "User", "הודעות"],
    ["POST", "/api/notifications/{id}/read", "Owner", "סימון נקראה"],
    ["GET", "/api/admin/dashboard", "Admin", "מדדים"],
    ["GET", "/api/admin/bookings", "Admin", "כל ההזמנות עם filters"],
    ["GET POST PUT", "/api/admin/resources", "Admin", "רשימה יצירה ועדכון"],
    ["GET POST", "/api/admin/resources/{id}/units", "Admin", "יחידות משאב"],
    ["PUT", "/api/admin/units/{id}", "Admin", "עדכון יחידה"],
    ["POST", "/api/admin/resources/{id}/sessions", "Admin", "יצירת Session"],
    ["PUT", "/api/admin/sessions/{id}", "Admin", "עדכון Session"],
    ["GET", "/api/admin/users", "Admin", "משתמשים"],
    ["GET", "/api/admin/audit", "Admin", "Audit"],
    ["GET", "/api/admin/system", "Admin", "Outbox והתראות"],
]
b.table(["Method", "Route", "הרשאה", "מטרה"], routes, font_size=7.8)

b.h1("סיכום מסלול לימוד")
b.p("כדי להבין את המערכת בקוד פתח תחילה Models.cs ו BookingRules.cs. לאחר מכן עבור ל HotelDbContext.cs כדי לראות כיצד האובייקטים הופכים לטבלאות. קרא את BookingService לפי הסדר Availability Hold Transition. המשך ל Controllers כדי לראות כיצד HTTP מגיע לשירותים. רק אז חזור ל Frontend: התחל ב types.ts ו api.ts עבור ל auth.tsx components.tsx App.tsx ולבסוף Pages.")
b.p("בשלב השני צור הזמנה בדפדפן והשאר את Developer Tools פתוח בלשונית Network. עבור כל בקשה מצא את Route ב Controller את המתודה בשירות את LINQ ב DbContext ואת השורה שנוצרה במסד. לאחר מכן עקוב אחר Outbox אל Worker Inbox Notification ו Elasticsearch. כאשר אתה מסוגל לספר את המסלול הזה בלי להסתכל בקוד אתה מבין את הארכיטקטורה של המערכת.")

b.save()
print(DOCX)
print(MD)
